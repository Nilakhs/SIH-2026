import os
import logging
import sqlite3
import json
from typing import Optional

logger = logging.getLogger(__name__)

async def process_document(doc_id: str, file_path: str, mime_type: str, db_path: str):
    """
    Process document in background.
    """
    try:
        conn = sqlite3.connect(db_path, check_same_thread=False)
        cursor = conn.cursor()
        
        # Update status to PROCESSING
        cursor.execute("UPDATE documents SET status = 'PROCESSING' WHERE id = ?", (doc_id,))
        conn.commit()

        page_count = 0
        chunks = []

        if mime_type == 'application/pdf' or file_path.lower().endswith('.pdf'):
            import fitz  # PyMuPDF
            try:
                import pytesseract
            except ImportError:
                pytesseract = None
            
            doc = fitz.open(file_path)
            page_count = len(doc)
            for page_num in range(page_count):
                page = doc.load_page(page_num)
                text = page.get_text()
                
                # If very little text, maybe it's scanned, try OCR
                if len(text.strip()) < 50:
                    image_list = page.get_images(full=True)
                    if image_list:
                        if pytesseract is not None:
                            try:
                                pix = page.get_pixmap()
                                temp_img = f"temp_{doc_id}_{page_num}.png"
                                pix.save(temp_img)
                                try:
                                    ocr_text = pytesseract.image_to_string(temp_img)
                                    text += "\n" + ocr_text
                                except pytesseract.TesseractNotFoundError:
                                    logger.warning("TesseractNotFoundError")
                                    text += "\n[OCR unavailable - Tesseract not installed]"
                                finally:
                                    if os.path.exists(temp_img):
                                        os.remove(temp_img)
                            except Exception as e:
                                logger.warning(f"OCR failed: {e}")
                                text += "\n[OCR failed]"
                        else:
                            text += "\n[OCR unavailable - Tesseract not installed]"
                
                meta = json.dumps({"page": page_num + 1})
                chunks.append((doc_id, page_num + 1, meta, text))
            doc.close()

        elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' or file_path.lower().endswith('.docx'):
            import docx
            doc = docx.Document(file_path)
            paragraphs = doc.paragraphs
            chunk_index = 0
            current_chunk = []
            start_para = 1
            
            for i, p in enumerate(paragraphs):
                current_chunk.append(p.text)
                if (i + 1) % 50 == 0:
                    chunk_index += 1
                    end_para = i + 1
                    meta = json.dumps({"paragraphs": f"{start_para}-{end_para}"})
                    chunks.append((doc_id, chunk_index, meta, "\n".join(current_chunk)))
                    current_chunk = []
                    start_para = i + 2
            
            if current_chunk:
                chunk_index += 1
                end_para = len(paragraphs)
                meta = json.dumps({"paragraphs": f"{start_para}-{end_para}"})
                chunks.append((doc_id, chunk_index, meta, "\n".join(current_chunk)))
            page_count = chunk_index

        elif file_path.lower().endswith(('.txt', '.csv')):
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                lines = f.readlines()
            
            chunk_index = 0
            current_chunk = []
            start_line = 1
            for i, line in enumerate(lines):
                current_chunk.append(line)
                if (i + 1) % 100 == 0:
                    chunk_index += 1
                    end_line = i + 1
                    meta = json.dumps({"lines": f"{start_line}-{end_line}"})
                    chunks.append((doc_id, chunk_index, meta, "".join(current_chunk)))
                    current_chunk = []
                    start_line = i + 2
            
            if current_chunk:
                chunk_index += 1
                end_line = len(lines)
                meta = json.dumps({"lines": f"{start_line}-{end_line}"})
                chunks.append((doc_id, chunk_index, meta, "".join(current_chunk)))
            page_count = chunk_index
        
        elif file_path.lower().endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            chunk_index = 0
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_text = []
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join([str(v) if v is not None else "" for v in row])
                    sheet_text.append(row_text)
                chunk_index += 1
                meta = json.dumps({"sheet": sheet, "rows": f"1-{len(sheet_text)}"})
                chunks.append((doc_id, chunk_index, meta, "\n".join(sheet_text)))
            page_count = chunk_index

        else:
            raise ValueError(f"Unsupported document format: {mime_type}")

        # Insert chunks
        cursor.executemany(
            "INSERT INTO document_chunks (document_id, chunk_index, metadata, content) VALUES (?, ?, ?, ?)",
            chunks
        )
        
        # Update to EXTRACTED
        cursor.execute("UPDATE documents SET status = 'EXTRACTED' WHERE id = ?", (doc_id,))
        conn.commit()
        
        # Update to INDEXING
        cursor.execute("UPDATE documents SET status = 'INDEXING' WHERE id = ?", (doc_id,))
        conn.commit()
        
        try:
            import sys
            import os
            sys.path.append(os.path.dirname(os.path.dirname(__file__)))
            from rag.embeddings import OllamaEmbeddingProvider
            from rag.vector_store import QdrantStore
            from rag.service import RAGService
            
            ep = OllamaEmbeddingProvider(host="http://localhost:11434")
            vs = QdrantStore(url="http://localhost:6333")
            rag_service = RAGService(ep, vs)
            
            rag_service.index_document(doc_id, db_path)
        except Exception as e:
            logger.error(f"Failed to index document {doc_id}: {e}")
            raise e

        # Update document status
        cursor.execute(
            "UPDATE documents SET status = 'COMPLETED', page_count = ? WHERE id = ?",
            (page_count, doc_id)
        )
        conn.commit()

    except Exception as e:
        logger.error(f"Failed to process document {doc_id}: {e}")
        try:
            conn = sqlite3.connect(db_path, check_same_thread=False)
            conn.execute(
                "UPDATE documents SET status = 'FAILED', error_msg = ? WHERE id = ?",
                (str(e), doc_id)
            )
            conn.commit()
        except Exception:
            pass
    finally:
        try:
            conn.close()
        except:
            pass
